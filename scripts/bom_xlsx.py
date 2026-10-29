#!/usr/bin/env python3
"""Build a Low/High Azure BOM workbook from a JSON spec. Formulas live, blue cells are inputs.

usage: bom_xlsx.py spec.json out.xlsx
Sheets: Summary (per block, Low/High, local currency, calculator links), Azure BOM (lines), Assumptions,
Other licences (optional), Sources. See reference/spec_example.json for the shape.
Monthly line cost = qty * unit_price * mult ; annual = monthly * 12 ; buffer applied on Summary total only.
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

spec = json.load(open(sys.argv[1])); OUT = sys.argv[2]
M = spec["meta"]; CUR = M.get("local_currency"); RATE = M.get("local_rate", 1)
NAVY, TINT = "3D2645", "F3E8F3"
th = Side(style="thin", color="E2DCE7"); BOX = Border(left=th, right=th, top=th, bottom=th)
HF = Font(bold=True, color="FFFFFF", name="Calibri", size=10); HFILL = PatternFill("solid", fgColor=NAVY)
INP = Font(color="0000FF", name="Calibri", size=10); NF = Font(name="Calibri", size=10)
BF = Font(bold=True, name="Calibri", size=10); TF = Font(bold=True, name="Calibri", size=14, color=NAVY)
NOTE = Font(italic=True, name="Calibri", size=9, color="6E6875")
USD = '$#,##0;[Red]-$#,##0'; USD4 = '$#,##0.0000'; LOC = '#,##0'
wb = Workbook()

def hdr(ws, r, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(r, i, c); x.font = HF; x.fill = HFILL; x.border = BOX
        x.alignment = Alignment(wrap_text=True, vertical="center")
def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[L(i)].width = x
def cell(ws, r, c, v, font=NF, fmt=None):
    x = ws.cell(r, c, v); x.font = font; x.border = BOX
    if fmt: x.number_format = fmt
    return x

# ---- Assumptions: A = name, B = Low, C = High, D = unit, E = basis. Row refs by key.
wa = wb.active; wa.title = "Assumptions"
wa["A1"] = f'{M["title"]} - costing assumptions'; wa["A1"].font = TF
wa["A2"] = "Blue = input. Change a blue cell and the BOM and Summary recalculate."; wa["A2"].font = NOTE
hdr(wa, 4, ["Input", "Low", "High", "Unit", "Basis"])
base = [("hours", "Hours per month", 730, 730, "hours", "8,760 / 12"),
        ("days", "Days per month", 30.4, 30.4, "days", ""),
        ("buffer", "Usage buffer on Azure total", M.get("buffer", 0.15), M.get("buffer", 0.15), "share", "Contingency for usage above plan; not in the calculator links")]
if CUR: base.append(("rate", f"{CUR} per USD", RATE, RATE, CUR, M.get("rate_basis", "")))
AREF = {}
for i, (k, n, lo, hi, u, b) in enumerate(base + [tuple(a) for a in spec.get("assumptions", [])]):
    r = 5 + i; AREF[k] = r
    cell(wa, r, 1, n); cell(wa, r, 2, lo, INP); cell(wa, r, 3, hi, INP); cell(wa, r, 4, u); cell(wa, r, 5, b)
widths(wa, [46, 12, 12, 14, 90])
A = lambda k, col: f"Assumptions!${col}${AREF[k]}"   # col B=Low, C=High

# ---- Azure BOM
wbm = wb.create_sheet("Azure BOM")
wbm["A1"] = f'{M["title"]} - Azure bill of materials'; wbm["A1"].font = TF
wbm["A2"] = f'Microsoft pay-as-you-go list prices, USD, {M["region"]}, as of {M["price_date"]} (Azure Retail Prices API / Azure Pricing Calculator). Excludes tax and any EA/CSP discount.'; wbm["A2"].font = NOTE
H = ["#", "Category", "Azure service / SKU", "Environment", "Region", "Qty Low", "Qty High", "Qty unit", "Unit price (USD)", "Price unit",
     "Multiplier / month", "Monthly Low (USD)", "Monthly High (USD)", "Annual Low (USD)", "Annual High (USD)", "Notes / source"]
hdr(wbm, 4, H); R0 = 5
for i, ln in enumerate(spec["lines"]):
    r = R0 + i
    def q(v, col):   # qty may be a number or an assumption key like "@users"
        return f"={A(v[1:], col)}" if isinstance(v, str) and v.startswith("@") else v
    mult = ln.get("mult", 1)
    if mult == "hours": mult = f"={A('hours','B')}"
    elif mult == "days": mult = f"={A('days','B')}"
    vals = [i + 1, ln["cat"], ln["item"], ln.get("env", "Prod"), ln.get("region", M["region"]), q(ln["low"], "B"), q(ln["high"], "C"),
            ln.get("qty_unit", ""), ln["price"], ln.get("price_unit", "per unit / month"), mult]
    for c, v in enumerate(vals, 1):
        cell(wbm, r, c, v, INP if c in (6, 7, 9, 11) else NF, USD4 if c == 9 else None)
    cell(wbm, r, 12, f"=F{r}*I{r}*K{r}", NF, USD); cell(wbm, r, 13, f"=G{r}*I{r}*K{r}", NF, USD)
    cell(wbm, r, 14, f"=L{r}*12", NF, USD); cell(wbm, r, 15, f"=M{r}*12", NF, USD)
    cell(wbm, r, 16, ln.get("note", ""))
R1 = R0 + len(spec["lines"]) - 1; RT = R1 + 1
cell(wbm, RT, 3, "Total Azure (before buffer)", BF)
for c in (12, 13, 14, 15): cell(wbm, RT, c, f"=SUM({L(c)}{R0}:{L(c)}{R1})", BF, USD)
widths(wbm, [4, 26, 58, 11, 14, 9, 9, 16, 13, 20, 12, 14, 14, 14, 14, 60])
wbm.freeze_panes = "D5"

# ---- Other licences (optional)
lic = spec.get("licences", [])
if lic:
    wl = wb.create_sheet("Other licences")
    wl["A1"] = "Other recurring licences and subscriptions (outside Azure)"; wl["A1"].font = TF
    hdr(wl, 3, ["#", "Item", "Qty Low", "Qty High", "Unit price (USD)", "Per", "Months / yr", "Annual Low (USD)", "Annual High (USD)", "Note"])
    for i, x in enumerate(lic):
        r = 4 + i
        for c, v in enumerate([i + 1, x["item"], x["low"], x["high"], x["price"], x.get("per", "unit-month"), x.get("months", 12)], 1):
            cell(wl, r, c, v, INP if c in (3, 4, 5, 7) else NF)
        cell(wl, r, 8, f"=C{r}*E{r}*G{r}", NF, USD); cell(wl, r, 9, f"=D{r}*E{r}*G{r}", NF, USD); cell(wl, r, 10, x.get("note", ""))
    LT = 4 + len(lic); cell(wl, LT, 2, "Total licences", BF)
    cell(wl, LT, 8, f"=SUM(H4:H{LT-1})", BF, USD); cell(wl, LT, 9, f"=SUM(I4:I{LT-1})", BF, USD)
    widths(wl, [4, 60, 9, 9, 14, 14, 10, 15, 15, 50])

# ---- Summary (first sheet)
ws = wb.create_sheet("Summary", 0)
ws["A1"] = f'{M["title"]} - recurring Azure cost, budgetary estimate'; ws["A1"].font = TF
ws["A2"] = M.get("disclaimer", "Non-binding budgetary estimate. Microsoft list prices, excluding tax and any negotiated discount. Low = planning minimum, High = planning ceiling.")
ws["A2"].font = NOTE
cols = ["#", "Block", "Low (USD / yr)", "High (USD / yr)"] + ([f"Low ({CUR} / yr)", f"High ({CUR} / yr)"] if CUR else []) + ["What is in it"]
hdr(ws, 4, cols)
cats = list(dict.fromkeys(l["cat"] for l in spec["lines"]))
rng = lambda c: f"'Azure BOM'!${c}${R0}:${c}${R1}"
for i, c in enumerate(cats):
    r = 5 + i
    cell(ws, r, 1, i + 1); cell(ws, r, 2, c.split(". ", 1)[-1])
    cell(ws, r, 3, f'=SUMIF({rng("B")},"{c}",{rng("N")})', NF, USD); cell(ws, r, 4, f'=SUMIF({rng("B")},"{c}",{rng("O")})', NF, USD)
    k = 5
    if CUR:
        cell(ws, r, 5, f"=C{r}*{A('rate','B')}", NF, LOC); cell(ws, r, 6, f"=D{r}*{A('rate','B')}", NF, LOC); k = 7
    cell(ws, r, k, spec.get("blocks", {}).get(c, ""))
rs = 5 + len(cats)
def line(r, label, lo, hi, font=NF):
    cell(ws, r, 2, label, font); cell(ws, r, 3, lo, font, USD); cell(ws, r, 4, hi, font, USD)
    if CUR: cell(ws, r, 5, f"=C{r}*{A('rate','B')}", font, LOC); cell(ws, r, 6, f"=D{r}*{A('rate','B')}", font, LOC)
line(rs, "Azure subtotal (before buffer)", f"=SUM(C5:C{rs-1})", f"=SUM(D5:D{rs-1})", BF)
line(rs + 1, "Usage buffer", f"=C{rs}*{A('buffer','B')}", f"=D{rs}*{A('buffer','C')}")
line(rs + 2, "Azure total incl. buffer", f"=C{rs}+C{rs+1}", f"=D{rs}+D{rs+1}", BF)
nxt = rs + 3
if lic:
    line(nxt, "Other licences (outside Azure)", f"='Other licences'!H{LT}", f"='Other licences'!I{LT}")
    line(nxt + 1, "Grand total per year", f"=C{rs+2}+C{nxt}", f"=D{rs+2}+D{nxt}", BF); nxt += 2
line(nxt + 1, "Azure monthly (incl. buffer)", f"=C{rs+2}/12", f"=D{rs+2}/12")
links = M.get("calculator", {})
if links:
    r = nxt + 3; cell(ws, r, 2, "Azure Pricing Calculator (list prices, excludes buffer)", BF)
    for j, (k, v) in enumerate(links.items()):
        cell(ws, r + 1 + j, 2, k); c = cell(ws, r + 1 + j, 3, v); c.hyperlink = v; c.font = Font(color="0563C1", underline="single", name="Calibri", size=10)
widths(ws, [4, 44, 16, 16, 16, 16, 90] if CUR else [4, 44, 16, 16, 90])

# ---- Sources
wsrc = wb.create_sheet("Sources")
wsrc["A1"] = "Sources and method"; wsrc["A1"].font = TF
src = ["Azure Retail Prices API (https://prices.azure.com/api/retail/prices), queried " + M["price_date"],
       "Azure Pricing Calculator (https://azure.microsoft.com/pricing/calculator/) for cross-check and share links",
       "Monthly cost = quantity x unit price x multiplier; 730 hours per month; annual = monthly x 12",
       "Usage buffer applied once on the Azure subtotal in Summary; calculator links exclude it"] + spec.get("sources", [])
for i, s in enumerate(src): wsrc.cell(3 + i, 1, f"{i+1}. {s}").font = NF
wsrc.column_dimensions["A"].width = 140
wb.save(OUT); print("wrote", OUT, "lines", len(spec["lines"]), "blocks", len(cats))
